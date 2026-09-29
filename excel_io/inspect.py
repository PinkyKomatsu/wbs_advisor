"""実用WBS様式（出力テンプレート）の解析。書き込み時の保護と、書き込み後の検証に使う。

記録するもの
  formulas        数式セル（シート全体の番地一覧）
  merged          結合セル
  validations     入力規則（範囲・種類・許可値）
  protected       シート保護の有無／ unlocked：ロック解除されたセル（保護時に書き込めるセル）
  names           名前定義
  vba             VBA プロジェクトの有無と vbaProject.bin の SHA-256
  shapes          図形・画像の数、forms：フォームコントロール（ボタン等）の数、activex：ActiveX の数
  ribbon          リボン設定（customUI）の有無
  cond_formats    条件付き書式の数、print：印刷設定
"""
from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl.utils import get_column_letter, range_boundaries

from . import reader

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def sheet_part(zf: zipfile.ZipFile, sheet: str) -> str:
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels.iter("{%s}Relationship" % NS_PKG)}
    for s in wb.iter("{%s}sheet" % NS_MAIN):
        if s.get("name") == sheet:
            t = targets[s.get("{%s}id" % NS_R)]
            return t.lstrip("/") if t.startswith("/") else "xl/" + t
    raise KeyError(f"シート「{sheet}」が見つかりません")


def _sheet_rels_targets(zf: zipfile.ZipFile, part: str) -> list[tuple[str, str]]:
    d, name = part.rsplit("/", 1)
    rels = f"{d}/_rels/{name}.rels"
    if rels not in zf.namelist():
        return []
    root = etree.fromstring(zf.read(rels))
    out = []
    for r in root.iter("{%s}Relationship" % NS_PKG):
        t = r.get("Target") or ""
        full = t.lstrip("/") if t.startswith("/") else re.sub(r"[^/]+/\.\./", "", f"{d}/{t}")
        out.append(((r.get("Type") or "").rsplit("/", 1)[-1], full))
    return out


def count_parts(zf: zipfile.ZipFile, part: str) -> dict:
    """シートに付いている図形・コントロールの数を数える。"""
    shapes = forms = activex = 0
    for typ, target in _sheet_rels_targets(zf, part):
        if target not in zf.namelist():
            continue
        data = zf.read(target)
        if typ == "drawing":
            shapes += len(re.findall(rb"<xdr:(?:sp|pic|graphicFrame|cxnSp)[ >]", data))
        elif typ == "vmlDrawing":
            forms += len(re.findall(rb'ObjectType="(?:Button|Checkbox|Drop|List|Radio|Spin|Scroll|GBox|Label|Edit)"', data))
        elif typ == "control":
            activex += 1
    if not forms:
        forms = len([n for n in zf.namelist() if n.startswith("xl/ctrlProps/")])
    return {"shapes": shapes, "forms": forms, "activex": activex}


def _allowed_values(wb, dv) -> list[str] | None:
    """リストの入力規則の許可値。求められなければ None。"""
    f = (dv.formula1 or "").strip()
    if not f:
        return None
    if f.startswith('"') and f.endswith('"'):
        return [x.strip() for x in f[1:-1].split(",")]
    ref = f.lstrip("=")
    if ref in wb.defined_names:
        ref = wb.defined_names[ref].attr_text
    m = re.match(r"^'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", ref)
    if not m:
        return None
    sheet = m.group(1)
    if sheet not in wb.sheetnames:
        return None
    rng = f"{m.group(2)}{m.group(3)}" + (f":{m.group(4)}{m.group(5)}" if m.group(4) else "")
    cells = wb[sheet][rng]
    if not isinstance(cells, tuple):
        cells = ((cells,),)
    vals = []
    for row in cells:
        for c in (row if isinstance(row, tuple) else (row,)):
            if c.value not in (None, ""):
                vals.append(reader.cell_text(c.value))
    return vals


NS_X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
NS_XM = "http://schemas.microsoft.com/office/excel/2006/main"


def _ext_validations(zf: zipfile.ZipFile, part: str) -> list[dict]:
    """拡張領域（x14:dataValidation）の入力規則。別シートを参照するリストは Excel がここに保存し、openpyxl は読まない。"""
    root = etree.fromstring(zf.read(part))
    out = []
    for dv in root.iter("{%s}dataValidation" % NS_X14):
        f1 = dv.find("{%s}formula1/{%s}f" % (NS_X14, NS_XM))
        sq = dv.find("{%s}sqref" % NS_XM)
        out.append({"type": dv.get("type"), "formula1": f1.text if f1 is not None else None,
                    "sqref": sq.text if sq is not None else ""})
    return out


class _DV:   # _allowed_values に渡すための入れ物
    def __init__(self, formula1):
        self.formula1 = formula1


def analyze(path, sheet: str, profile: dict | None = None) -> dict:
    p = reader.check(path)
    wb = reader.load(p, data_only=False)
    ws = wb[sheet]
    formulas = []
    unlocked = []
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                formulas.append(c.coordinate)
            try:
                if c.protection is not None and c.protection.locked is False:
                    unlocked.append(c.coordinate)
            except AttributeError:
                pass
    validations = []
    for dv in ws.data_validations.dataValidation:
        validations.append({
            "sqref": str(dv.sqref), "type": dv.type, "formula1": dv.formula1,
            "allowed": _allowed_values(wb, dv) if dv.type == "list" else None,
        })
    names = sorted(str(n) for n in wb.defined_names)
    for w in wb.worksheets:
        for n in getattr(w, "defined_names", {}) or {}:
            names.append(f"{w.title}!{n}")
    with zipfile.ZipFile(p) as zf:
        part = sheet_part(zf, sheet)
        for x in _ext_validations(zf, part):
            validations.append({"sqref": x["sqref"], "type": x["type"], "formula1": x["formula1"], "ext": True,
                                "allowed": _allowed_values(wb, _DV(x["formula1"])) if x["type"] == "list" else None})
        parts = count_parts(zf, part)
        vba = "xl/vbaProject.bin" in zf.namelist()
        vba_hash = hashlib.sha256(zf.read("xl/vbaProject.bin")).hexdigest() if vba else None
        ribbon = any(n.lower().startswith("customui") for n in zf.namelist())
    col_widths = {k: v.width for k, v in ws.column_dimensions.items() if v.width}
    row_heights = {str(k): v.height for k, v in ws.row_dimensions.items() if v.height}
    return {
        "sheet": sheet,
        "part": part,
        "formulas": formulas,
        "merged": sorted(str(r) for r in ws.merged_cells.ranges),
        "validations": validations,
        "protected": bool(ws.protection.sheet),
        "unlocked": unlocked,
        "names": sorted(set(names)),
        "vba": vba,
        "vba_sha256": vba_hash,
        "ribbon": ribbon,
        **parts,
        "cond_formats": sum(len(cf.rules) for cf in ws.conditional_formatting),
        # ルールの中身（行を挿入すると Excel が同じルールを範囲ごとに分けるため、件数ではなく中身で比べる）
        "cond_rules": sorted({f"{r.type}|{r.operator}|{'|'.join(r.formula or [])}"
                              for cf in ws.conditional_formatting for r in cf.rules}),
        "print": {"area": ws.print_area, "titles": ws.print_title_rows, "orientation": ws.page_setup.orientation},
        "col_widths": col_widths,
        "row_heights": row_heights,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
    }


def summary_lines(a: dict) -> list[str]:
    """様式登録画面に出す説明。"""
    return [
        f"数式セル：{len(a['formulas'])}個",
        f"結合セル：{len(a['merged'])}箇所",
        f"入力規則：{len(a['validations'])}件" + ("".join(
            f"\n  ・{v['sqref']}：{'、'.join(v['allowed'])}" for v in a["validations"] if v.get("allowed"))),
        f"シート保護：{'あり' if a['protected'] else 'なし'}" + (f"（ロック解除セル {len(a['unlocked'])}個）" if a["protected"] else ""),
        f"名前定義：{len(a['names'])}件（{'、'.join(a['names'][:5])}{'…' if len(a['names']) > 5 else ''}）",
        f"VBA プロジェクト：{'あり' if a['vba'] else 'なし'}",
        f"図形・画像：{a['shapes']}個　フォームコントロール：{a['forms']}個　ActiveX：{a['activex']}個",
        f"リボン設定：{'あり' if a['ribbon'] else 'なし'}　条件付き書式：{a['cond_formats']}件",
    ]


def validation_for(analysis: dict, ref: str) -> dict | None:
    """セルにかかる入力規則。"""
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    col, row = coordinate_from_string(ref)
    c = column_index_from_string(col)
    for v in analysis.get("validations", []):
        for part in v["sqref"].split():
            min_c, min_r, max_c, max_r = range_boundaries(part)
            if min_c <= c <= max_c and min_r <= row <= max_r:
                return v
    return None


def merged_anchor(analysis: dict, ref: str) -> str | None:
    """結合セルの一部なら左上のセル（左上そのものなら自身）。結合されていなければ None。"""
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    col, row = coordinate_from_string(ref)
    c = column_index_from_string(col)
    for m in analysis.get("merged", []):
        min_c, min_r, max_c, max_r = range_boundaries(m)
        if min_c <= c <= max_c and min_r <= row <= max_r:
            return f"{get_column_letter(min_c)}{min_r}"
    return None
