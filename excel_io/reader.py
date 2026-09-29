"""Excel の読込（openpyxl。保存はしない）。

- read_only=True：大量の行を速く読む（ひな形・過去事例の明細）
- 通常モード：結合セル・入力規則・保護・書式を読む（プロファイル推定・様式解析）
  図形や画像は openpyxl が Pillow を必要とするため、シートの rels から drawing の参照を外したコピーを
  メモリ上で読む（元のファイルは変更しない）。
"""
from __future__ import annotations

import hashlib
import io
import re
import warnings
import zipfile
from pathlib import Path

import openpyxl
from lxml import etree

SUPPORTED = (".xlsx", ".xlsm")


class UnsupportedFile(Exception):
    pass


def check(path) -> Path:
    p = Path(path)
    if p.suffix.lower() == ".xls":
        raise UnsupportedFile(f"「{p.name}」は旧形式（.xls）です。Excel で .xlsx または .xlsm として保存し直してください。")
    if p.suffix.lower() not in SUPPORTED:
        raise UnsupportedFile(f"「{p.name}」は対応していない形式です（.xlsx / .xlsm のみ）。")
    if not zipfile.is_zipfile(p):
        raise UnsupportedFile(f"「{p.name}」を Excel ファイルとして開けません。")
    return p


def sha256(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _without_drawings(path: Path) -> io.BytesIO:
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if re.match(r"^xl/worksheets/_rels/[^/]+\.rels$", info.filename):
                root = etree.fromstring(data)
                for r in list(root):
                    if (r.get("Type") or "").endswith("/drawing"):
                        root.remove(r)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info.filename, data)
    buf.seek(0)
    return buf


def load(path, data_only: bool = True, read_only: bool = False):
    p = check(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        if read_only:
            return openpyxl.load_workbook(p, read_only=True, data_only=data_only)
        return openpyxl.load_workbook(_without_drawings(p), data_only=data_only)


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def iter_rows(path, sheet: str, start_row: int, columns: dict[str, str], blank_stop: int = 5):
    """start_row から、列の意味ごとの値を行単位で返す（空行が blank_stop 行続いたら終わり）。
    columns: {"B": "wbs_no", "D": "name", ...}"""
    from openpyxl.utils import column_index_from_string
    wb = load(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet]
        idx = {column_index_from_string(c) - 1: m for c, m in columns.items()}
        blanks = 0
        for r, row in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start_row):
            values = {m: row[i] if i < len(row) else None for i, m in idx.items()}
            if all(v in (None, "") for v in values.values()):
                blanks += 1
                if blanks >= blank_stop:
                    break
                continue
            blanks = 0
            yield r, values
    finally:
        wb.close()
