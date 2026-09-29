"""書き込み後の検証：出力ファイルとテンプレートを比べる。

OOXML 方式：対象シート XML 以外の全パーツがバイト単位で一致すること、シート XML も明細（sheetData）以外は一致し、
            書き込み対象以外のセルは値もスタイルも変わっていないこと。
COM 方式  ：Excel が保存し直すためパーツのバイト一致は求めず、構造で比べる
            （vbaProject.bin・図形・フォームコントロール・ActiveX・名前定義・結合・入力規則・条件付き書式・
             シート保護・印刷設定・列幅・行高・書き込み対象以外のセルの書式）。
"""
from __future__ import annotations

import hashlib
import zipfile
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from openpyxl.utils import range_boundaries

from . import inspect, reader

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
M = "{%s}" % NS


@dataclass
class Check:
    name: str
    ok: bool
    detail: str = ""


def _c14n(el) -> bytes:
    return etree.tostring(el, method="c14n")


def _rgb(color) -> str | None:
    rgb = getattr(color, "rgb", None) if color is not None else None
    return rgb[-6:].upper() if isinstance(rgb, str) and len(rgb) >= 6 else None   # 透明度（先頭 2 桁）は無視


def same_style(a, b) -> bool:
    """実際に効いている書式が同じか（Excel が保存し直したときの表し方の違いは同じとみなす）。
    - テーマのフォント（scheme 指定）や未指定のフォント名・サイズは、比べない（表示は同じフォント）
    - 色は透明度を除いた RGB で比べる"""
    fa, fb = a.font, b.font
    if fa.name is not None and not fa.scheme and fa.name != fb.name:
        return False
    if fa.sz is not None and fa.sz != fb.sz:
        return False
    if bool(fa.b) != bool(fb.b) or bool(fa.i) != bool(fb.i):
        return False
    ca, cb = _rgb(fa.color), _rgb(fb.color)
    if ca and cb and ca != cb:
        return False
    for s in ("left", "right", "top", "bottom"):
        if getattr(getattr(a.border, s), "style", None) != getattr(getattr(b.border, s), "style", None):
            return False
    if (a.fill.fill_type or None) != (b.fill.fill_type or None):
        return False
    if a.fill.fill_type == "solid" and _rgb(a.fill.fgColor) != _rgb(b.fill.fgColor):
        return False
    if a.number_format != b.number_format:
        return False
    return (a.alignment.horizontal or None) == (b.alignment.horizontal or None) and \
        bool(a.alignment.wrap_text) == bool(b.alignment.wrap_text)


def col_widths(path, sheet: str) -> dict[str, float]:
    """列ごとの幅（Excel が同じ幅の列を「C〜D」とまとめて保存した場合も列ごとに展開する）。"""
    from openpyxl.utils import get_column_letter
    ws = reader.load(path, data_only=False)[sheet]
    out = {}
    for dim in ws.column_dimensions.values():
        if dim.width and dim.min and dim.max:
            for c in range(dim.min, dim.max + 1):
                out[get_column_letter(c)] = round(float(dim.width), 2)
    return out


def _style_sig(cell) -> tuple:
    f, b, fill = cell.font, cell.border, cell.fill
    return (f.name, f.sz, f.b, f.i, f.color.rgb if f.color is not None else None,
            tuple(getattr(getattr(b, s), "style", None) for s in ("left", "right", "top", "bottom")),
            fill.fill_type, fill.fgColor.rgb if fill.fgColor is not None else None,
            cell.number_format, cell.alignment.horizontal, cell.alignment.wrap_text)


def _shift_ref(ref: str, at: int | None, n: int) -> str:
    """行挿入前の番地を、挿入後の番地に直す。"""
    if not n or at is None:
        return ref
    min_c, min_r, max_c, max_r = range_boundaries(ref)
    from openpyxl.utils import get_column_letter as L
    r1 = min_r + n if min_r >= at else min_r
    r2 = max_r + n if max_r >= at else max_r
    a = f"{L(min_c)}{r1}"
    return a if ref.count(":") == 0 else f"{a}:{L(max_c)}{r2}"


def verify(template, out, plan, analysis: dict) -> list[Check]:
    template, out = Path(template), Path(out)
    checks: list[Check] = []
    with zipfile.ZipFile(template) as a, zipfile.ZipFile(out) as b:
        names_a, names_b = set(a.namelist()), set(b.namelist())
        part = inspect.sheet_part(a, plan.sheet)
        # VBA
        if "xl/vbaProject.bin" in names_a:
            same = "xl/vbaProject.bin" in names_b and a.read("xl/vbaProject.bin") == b.read("xl/vbaProject.bin")
            checks.append(Check("VBA プロジェクト（vbaProject.bin）", same,
                                "一致" if same else "内容が変わっています"))
        if plan.engine == "ooxml":
            diff = [n for n in names_a | names_b if n != part and
                    (n not in names_a or n not in names_b or a.read(n) != b.read(n))]
            checks.append(Check("対象シート以外のパーツ（図形・スタイル等）", not diff,
                                "すべてバイト単位で一致" if not diff else "変わったパーツ：" + "、".join(sorted(diff))))
            sa, sb = etree.fromstring(a.read(part)), etree.fromstring(b.read(part))
            da, db_ = sa.find(M + "sheetData"), sb.find(M + "sheetData")
            sa.remove(da)
            sb.remove(db_)
            same = _c14n(sa) == _c14n(sb)
            checks.append(Check("シートの明細以外（列幅・結合・入力規則・保護・印刷設定）", same, "一致" if same else "差異あり"))
            targets = set(plan.cell_values())
            cells_a = {c.get("r"): c for c in da.iter(M + "c")}
            cells_b = {c.get("r"): c for c in db_.iter(M + "c")}
            changed = [r for r, c in cells_a.items() if r not in targets and (r not in cells_b or _c14n(c) != _c14n(cells_b[r]))]
            style_changed = [r for r in targets if r in cells_a and r in cells_b and cells_a[r].get("s") != cells_b[r].get("s")]
            checks.append(Check("書き込み対象以外のセル", not changed, "変化なし" if not changed else "変化：" + "、".join(changed[:10])))
            checks.append(Check("書き込んだセルの書式（スタイル番号）", not style_changed,
                                "変化なし" if not style_changed else "変化：" + "、".join(style_changed[:10])))
        else:
            pa, pb = inspect.count_parts(a, part), inspect.count_parts(b, inspect.sheet_part(b, plan.sheet))
            for key, label in (("shapes", "図形・画像"), ("forms", "フォームコントロール（ボタン等）"), ("activex", "ActiveX")):
                checks.append(Check(label, pa[key] == pb[key], f"{pa[key]}個 → {pb[key]}個"))
            ribbon_a = any(n.lower().startswith("customui") for n in names_a)
            ribbon_b = any(n.lower().startswith("customui") for n in names_b)
            if ribbon_a:
                checks.append(Check("リボン設定（customUI）", ribbon_b, "保持" if ribbon_b else "消えています"))

    # 構造（両方式共通）
    after = inspect.analyze(out, plan.sheet)
    n, at = plan.insert_rows, plan.insert_at
    names_ok = set(after["names"]) == set(analysis["names"])
    checks.append(Check("名前定義", names_ok, f"{len(analysis['names'])}件 → {len(after['names'])}件"))
    expected_merges = {_shift_ref(m, at, n) for m in analysis["merged"]}
    row_merges = [m for m in analysis["merged"] if range_boundaries(m)[1] == range_boundaries(m)[3] == (plan.source_row or -1)]
    merges_ok = expected_merges <= set(after["merged"]) and len(after["merged"]) == len(analysis["merged"]) + n * len(row_merges)
    checks.append(Check("結合セル", merges_ok, f"{len(analysis['merged'])}箇所 → {len(after['merged'])}箇所"
                        + (f"（挿入行の分 +{n * len(row_merges)}）" if n and row_merges else "")))
    val_ok = len(after["validations"]) == len(analysis["validations"]) and all(
        (va.get("allowed") == vb.get("allowed") and va.get("type") == vb.get("type"))
        for va, vb in zip(sorted(analysis["validations"], key=lambda v: v["formula1"] or ""),
                          sorted(after["validations"], key=lambda v: v["formula1"] or "")))
    checks.append(Check("入力規則", val_ok, f"{len(analysis['validations'])}件 → {len(after['validations'])}件"))
    same_rules = after.get("cond_rules") == analysis.get("cond_rules")
    count_ok = after["cond_formats"] == analysis["cond_formats"] or (n > 0 and after["cond_formats"] >= analysis["cond_formats"])
    checks.append(Check("条件付き書式", same_rules and count_ok,
                        f"ルール {len(analysis.get('cond_rules', []))}種類（{analysis['cond_formats']}件 → {after['cond_formats']}件"
                        + ("。挿入行の範囲に同じルールを適用" if after["cond_formats"] > analysis["cond_formats"] else "") + "）"))
    checks.append(Check("シート保護", after["protected"] == analysis["protected"], "あり" if after["protected"] else "なし"))
    pa_, pb_ = analysis["print"], after["print"]
    orient = lambda o: o or "portrait"            # 未指定は既定の「縦」
    print_ok = (pa_["titles"] or None) == (pb_["titles"] or None) and orient(pa_["orientation"]) == orient(pb_["orientation"]) \
        and ((pa_["area"] or None) == (pb_["area"] or None) or n > 0)
    checks.append(Check("印刷設定", print_ok, f"印刷範囲 {pa_['area'] or 'なし'} → {pb_['area'] or 'なし'}、"
                                            f"向き {orient(pa_['orientation'])} → {orient(pb_['orientation'])}"))
    wa_, wb2 = col_widths(template, plan.sheet), col_widths(out, plan.sheet)
    diff_w = [c for c in wa_ if wa_[c] != wb2.get(c)]
    checks.append(Check("列幅", not diff_w, "一致" if not diff_w else "差異：" + "、".join(diff_w)))
    if not n:
        checks.append(Check("行高", after["row_heights"] == analysis["row_heights"], "一致" if after["row_heights"] == analysis["row_heights"] else "差異あり"))

    # 書式（書き込み対象以外のセル）と書き込み結果
    wa, wb_ = reader.load(template, data_only=False), reader.load(out, data_only=False)
    ws_a, ws_b = wa[plan.sheet], wb_[plan.sheet]
    targets = set(plan.cell_values())
    diffs = []
    for row in ws_a.iter_rows(min_row=1, max_row=ws_a.max_row):
        for c in row:
            if c.coordinate in targets:
                continue
            ref = _shift_ref(c.coordinate, at, n)
            if not same_style(c, ws_b[ref]):
                diffs.append(c.coordinate)
    checks.append(Check("書式（フォント・罫線・塗り・表示形式）", not diffs,
                        "変化なし" if not diffs else "変化：" + "、".join(diffs[:10])))
    wrong = []
    for w in plan.actions:
        got = ws_b[w.ref].value
        want = None if w.status == "clear" else w.value
        if (got in (None, "") and want in (None, "")) or str(got) == str(want) or got == want:
            continue
        wrong.append(f"{w.ref}（期待 {want}／実際 {got}）")
    checks.append(Check("書き込んだ値", not wrong, f"{len(plan.actions)}セル" if not wrong else "違い：" + "、".join(wrong[:5])))
    return checks


def file_hash(path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()
