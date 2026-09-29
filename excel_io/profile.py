"""様式プロファイルの自動推定。

プロファイル（dict）
  sheet        対象シート
  header_row   ヘッダー行
  data_start   データ開始行
  data_end     データ終了行（実用WBS様式のみ。テンプレートの明細の最終行。None なら未確定）
  columns      {"B": "wbs_no", "D": "name", ...}（列 → 列の意味）
  headers      {"B": "WBS番号", ...}（列 → 見出しの文字列）
  confidence   0〜1
  issues       迷った点の説明（確信度が閾値未満ならプレビューでユーザーが指定する）
"""
from __future__ import annotations

from openpyxl.utils import get_column_letter

from core.normalize import norm

from . import reader

# 列の意味と見出しのキーワード（前にあるほど優先）
KEYWORDS = {
    "wbs_no": ["wbs番号", "wbsno", "wbs", "作業番号", "タスク番号", "項番"],
    "no": ["no", "№", "#", "連番", "番号"],
    "phase": ["工程", "フェーズ", "phase", "大項目", "区分", "カテゴリ"],
    "name": ["作業項目", "作業名", "作業内容", "タスク名", "タスク", "項目名", "項目", "作業", "task"],
    "owner": ["担当者", "担当", "責任者", "owner"],
    "effort": ["予定工数", "工数", "人日", "見積", "effort"],
    "start": ["開始予定", "開始日", "開始", "着手", "start"],
    "end": ["終了予定", "終了日", "終了", "完了予定", "完了", "end", "期限"],
    "status": ["状態", "ステータス", "進捗", "status"],
    "note": ["備考", "メモ", "コメント", "note", "remarks"],
}
REQUIRED = ("name",)
_TOTAL_WORDS = ("合計", "小計", "総計", "total")


def keywords(cfg: dict | None) -> dict[str, list[str]]:
    """設定の header_keywords（例：管理単位（中２）→ 大分類）を、組み込みのキーワードより優先して使う。"""
    extra = {m: [norm(w) for w in ws] for m, ws in (cfg or {}).get("header_keywords", {}).items()
             if isinstance(ws, list)}
    out = {m: extra.get(m, []) + [w for w in KEYWORDS.get(m, []) if w not in extra.get(m, [])]
           for m in list(extra) + [m for m in KEYWORDS if m not in extra]}
    return out


def _match(text: str, kw: dict[str, list[str]] | None = None) -> tuple[str | None, float]:
    t = norm(text)
    if not t or len(t) > 20:
        return None, 0.0
    best, best_s = None, 0.0
    for meaning, words in (kw or KEYWORDS).items():
        for rank, w in enumerate(words):
            s = 1.0 if t == w else (0.8 if w in t and len(w) >= 2 else 0.0)
            s -= rank * 0.01
            if s > best_s:
                best, best_s = meaning, s
    return best, best_s


def _header_candidates(ws, scan_rows: int, kw=None):
    cands = []
    for r in range(1, min(ws.max_row, scan_rows) + 1):
        found: dict[str, tuple[str, float, str]] = {}
        dup = []
        for c in range(1, min(ws.max_column, 60) + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            meaning, s = _match(v, kw)
            if not meaning:
                continue
            col = get_column_letter(c)
            if meaning in found:
                dup.append(meaning)
                if s <= found[meaning][1]:
                    continue
            found[meaning] = (col, s, v.strip())
        score = sum(s for _, s, _ in found.values()) + (1.0 if "name" in found else 0.0)
        if found:
            cands.append({"row": r, "found": found, "score": score, "dup": dup})
    return sorted(cands, key=lambda c: -c["score"])


def _is_blank_row(ws, r: int, cols: list[int]) -> bool:
    return all(ws.cell(r, c).value in (None, "") for c in cols)


def _has_border(cell) -> bool:
    b = cell.border
    return any(getattr(getattr(b, s, None), "style", None) for s in ("left", "right", "top", "bottom"))


def detect(path, kind: str, cfg: dict | None = None) -> dict:
    """kind: "master"（ひな形・過去事例） / "output"（実用WBS様式）。"""
    cfg = cfg or {}
    pcfg = cfg.get("profile", {})
    scan = int(pcfg.get("header_scan_rows", 30))
    kw = keywords(cfg)
    wb = reader.load(path, data_only=False)
    best = None
    ranking = []
    for ws in wb.worksheets:
        for cand in _header_candidates(ws, scan, kw)[:3]:
            ranking.append((cand["score"], ws.title, cand))
    ranking.sort(key=lambda x: -x[0])
    issues = []
    if not ranking:
        return {"sheet": wb.worksheets[0].title, "header_row": None, "data_start": None, "data_end": None,
                "columns": {}, "headers": {}, "confidence": 0.0, "sheets": wb.sheetnames,
                "issues": ["ヘッダー行が見つかりません。プレビューでヘッダー行と各列を指定してください。"]}
    score, sheet, best = ranking[0]
    ws = wb[sheet]
    found = best["found"]
    columns = {col: m for m, (col, _, _) in found.items()}
    headers = {col: text for m, (col, _, text) in found.items()}

    conf = 0.0
    if "name" in found:
        conf += 0.5
    else:
        issues.append("作業項目名の列が見つかりません。")
    if "wbs_no" in found:
        conf += 0.2
    conf += min(3, len([m for m in found if m not in ("name", "wbs_no")])) * 0.1
    if len(ranking) > 1 and ranking[1][0] >= score - 0.3 and (ranking[1][1], ranking[1][2]["row"]) != (sheet, best["row"]):
        conf -= 0.2
        issues.append(f"ヘッダー行の候補が複数あります（{ranking[1][1]} の {ranking[1][2]['row']} 行目）。")
    for m in best["dup"]:
        conf -= 0.1
        issues.append(f"「{m}」に当てはまる列が複数あります。")

    header_row = best["row"]
    data_start = header_row + 1
    col_idx = [ws[f"{c}1"].column for c in columns]
    # 見出しの下の空白行（2 段見出しの隙間など）は飛ばす（ひな形・過去事例のみ）
    if kind == "master":
        while data_start < header_row + 4 and _is_blank_row(ws, data_start, col_idx):
            data_start += 1

    data_end = None
    if kind == "output":
        name_col = next((c for c, m in columns.items() if m == "name"), None)
        r = data_start
        while name_col and r <= max(ws.max_row, data_start):
            row_texts = [str(ws.cell(r, c).value or "") for c in col_idx]
            if any(any(w in norm(t) for w in _TOTAL_WORDS) for t in row_texts):
                break
            cell = ws[f"{name_col}{r}"]
            if not _has_border(cell) and cell.value in (None, ""):
                break
            r += 1
        data_end = r - 1 if r > data_start else None
        if data_end is None:
            conf -= 0.2
            issues.append("テンプレートの明細行（空き行）が判別できません。プレビューで最終行を指定してください。")
    conf = round(max(0.0, min(1.0, conf)), 2)
    return {"sheet": sheet, "header_row": header_row, "data_start": data_start, "data_end": data_end,
            "columns": dict(sorted(columns.items(), key=lambda kv: (len(kv[0]), kv[0]))),
            "headers": headers, "confidence": conf, "issues": issues, "sheets": wb.sheetnames}


def needs_review(profile: dict, cfg: dict) -> bool:
    th = float(cfg.get("profile", {}).get("confidence_threshold", 0.7))
    return profile.get("confidence", 0) < th or bool(profile.get("issues"))


def read_items(path, profile: dict, blank_stop: int = 5) -> list[tuple[int, dict]]:
    """プロファイルに従って明細を読む。[(行番号, {意味: 値})]"""
    if not profile.get("columns") or not profile.get("data_start"):
        return []
    end = profile.get("data_end")
    out = []
    for r, values in reader.iter_rows(path, profile["sheet"], profile["data_start"], profile["columns"], blank_stop):
        if end and r > end:
            break
        if any(any(w in norm(reader.cell_text(v)) for w in _TOTAL_WORDS) for v in values.values()):
            continue
        out.append((r, values))
    return out
