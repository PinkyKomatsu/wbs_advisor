"""書き込み計画：どのセルに何を書くか（差分プレビューと、COM / OOXML 両方式の共通ルール）。

共通ルール
  - 結合セルは左上セルにのみ書き込む（左上以外に当たる値は左上へ。左上が別の列の書き込み先なら書かない）
  - 入力規則（リスト）のあるセルには許可値のみ書き込む。許可値でなければ値の対応表で置き換え、
    それでも許可値にならなければ書き込まずに警告する
  - 数式セル・（シート保護時の）ロックされたセルには書き込まない
  - 書式は変更しない（値だけを書く）
  - テンプレートの空き行が足りないとき：COM 方式は行を挿入する。OOXML 方式は書き込まずに不足行数を警告する
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl.utils import column_index_from_string, get_column_letter

from core.models import MEANINGS

from . import inspect, reader


CONST_PREFIX = "__const__:"   # 列に固定値を書く書き込み元
TYPE_SOURCE = "__type__"       # 列に WBS の種類（設計・構築・テスト）を書く書き込み元


@dataclass
class CellWrite:
    ref: str
    row: int
    col: str
    value: object
    meaning: str
    status: str            # write（書く） / clear（見本の値を消す） / skip（書かない）
    reason: str = ""
    current: str = ""


@dataclass
class WritePlan:
    sheet: str
    engine: str
    data_start: int
    data_end: int | None
    capacity: int
    rows_needed: int
    insert_rows: int = 0
    insert_at: int | None = None
    source_row: int | None = None
    writes: list[CellWrite] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    blocked: str | None = None

    @property
    def actions(self) -> list[CellWrite]:
        return [w for w in self.writes if w.status in ("write", "clear")]

    @property
    def skipped(self) -> list[CellWrite]:
        return [w for w in self.writes if w.status == "skip"]

    def cell_values(self) -> dict[str, object]:
        """書き込みエンジンに渡す {セル: 値}（clear は None）。"""
        return {w.ref: (None if w.status == "clear" else w.value) for w in self.actions}


def text_of(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%Y/%m/%d")
    if isinstance(v, date):
        return v.strftime("%Y/%m/%d")
    return reader.cell_text(v)


def template_values(path, sheet: str, first_row: int, last_row: int, cols: list[str]) -> dict[str, str]:
    """差分プレビュー用に、テンプレートの現在の値（数式はそのまま）を読む。"""
    wb = reader.load(path, data_only=False)
    ws = wb[sheet]
    out = {}
    for r in range(first_row, last_row + 1):
        for c in cols:
            v = ws[f"{c}{r}"].value
            if v not in (None, ""):
                out[f"{c}{r}"] = text_of(v)
    return out


def default_sources(template_profile: dict, master_meanings: set[str], analysis: dict | None = None) -> dict[str, str | None]:
    """テンプレートの各列に、ひな形のどの意味の値を入れるか（同じ意味の列があればそれ）。
    明細に数式が入っている列（例：No＝ROW()）は書き込まない。"""
    out = {}
    formulas = set((analysis or {}).get("formulas", []))
    start = template_profile.get("data_start") or 0
    end = template_profile.get("data_end") or start
    for col, meaning in template_profile.get("columns", {}).items():
        if any(f"{col}{r}" in formulas for r in range(start, end + 1)):
            out[col] = None
        elif meaning == "no":
            out[col] = "__seq__"         # 連番
        elif meaning in master_meanings or meaning in ("wbs_no", "phase", "name"):
            out[col] = meaning
        else:
            out[col] = None              # 対応が判断できない → 画面で選ぶ
    return out


def build_plan(analysis: dict, profile: dict, rows: list[dict], sources: dict[str, str | None],
               value_map: dict, engine: str, clear_sample: bool = True,
               current: dict[str, str] | None = None) -> WritePlan:
    sheet = profile["sheet"]
    start, end = profile.get("data_start"), profile.get("data_end")
    if not start:
        return WritePlan(sheet, engine, 0, None, 0, len(rows), blocked="データ開始行が決まっていません。プレビューで指定してください。")
    capacity = (end - start + 1) if end else 0
    plan = WritePlan(sheet, engine, start, end, capacity, len(rows))
    current = current or {}
    formulas = set(analysis.get("formulas", []))
    unlocked = set(analysis.get("unlocked", []))
    protected = bool(analysis.get("protected"))
    template_row = end or start                         # 挿入した行は、この行と同じ構成とみなす
    cols = [c for c, s in sources.items() if s]

    if len(rows) > capacity:
        short = len(rows) - capacity
        if engine == "ooxml":
            plan.blocked = (f"テンプレートの空き行が {short} 行不足しています（明細 {capacity} 行に対し {len(rows)} 行）。"
                            "OOXML 方式は行の挿入に対応していないため、書き込みを中止しました。"
                            "Excel がインストールされた PC で出力するか、テンプレートの空き行を増やしてください。")
            return plan
        if protected:
            plan.blocked = f"シートが保護されているため行を挿入できません（{short} 行不足）。テンプレートの保護を解除するか、空き行を増やしてください。"
            return plan
        plan.insert_rows = short
        plan.insert_at = end if end else start           # 明細の最終行の上に入れる（合計などの範囲が広がるように）
        plan.source_row = (end - 1) if end and end - 1 >= start else (end or start)
        plan.warnings.append(f"テンプレートの空き行が {short} 行不足するため、{short} 行を挿入します（直上行の書式・入力規則・数式をコピー）。")

    def geometry_ref(col: str, r: int) -> str:
        """保護・数式・結合・入力規則を調べるときの、テンプレート上の対応セル。"""
        if end and r > end:                   # 挿入後に下へずれた行・挿入した行
            return f"{col}{template_row}"
        return f"{col}{r}"

    def formula_col(col: str) -> bool:
        return any(f"{col}{r}" in formulas for r in range(start, (end or start) + 1))

    written_refs: set[str] = set()
    for i, values in enumerate(rows):
        r = start + i
        for col, src in sources.items():
            if not src:
                continue
            meaning = profile["columns"].get(col, "")
            if src == "__seq__":
                value = i + 1
            elif src.startswith(CONST_PREFIX):
                value = src[len(CONST_PREFIX):]          # 固定値（例：管理単位（大）に案件名）
            else:
                value = values.get(src)                   # __type__ なら WBS の種類
            if value in (None, ""):
                continue
            ref = f"{col}{r}"
            g = geometry_ref(col, r)
            w = CellWrite(ref, r, col, value, meaning, "write", current=current.get(ref, ""))
            if g in formulas or (plan.insert_rows and r >= (plan.insert_at or 0) and formula_col(col)):
                w.status, w.reason = "skip", "数式セル"
            elif protected and g not in unlocked:
                w.status, w.reason = "skip", "ロックされたセル（シート保護）"
            else:
                anchor = inspect.merged_anchor(analysis, g)
                if anchor and anchor != g:
                    a_col = re.match(r"[A-Z]+", anchor).group(0)
                    if sources.get(a_col):
                        w.status, w.reason = "skip", f"結合セルの左上以外（左上 {anchor}）"
                    else:
                        w.ref, w.col = f"{a_col}{r}", a_col
                        w.reason = "結合セルの左上に書き込み"
                if w.status == "write":
                    v = inspect.validation_for(analysis, geometry_ref(w.col, r))
                    if v and v.get("type") == "list":
                        allowed = v.get("allowed")
                        text = text_of(w.value)
                        if allowed is None:
                            w.status, w.reason = "skip", "入力規則の許可値を特定できません"
                        elif text not in allowed:
                            mapped = (value_map.get(meaning, {}) or {}).get(text) or (value_map.get("*", {}) or {}).get(text)
                            if mapped in allowed:
                                w.value, w.reason = mapped, f"値の対応表で置換（{text}→{mapped}）"
                            else:
                                w.status = "skip"
                                w.reason = f"入力規則の許可値にない値（{text}。許可値：{'、'.join(allowed)}）"
            if w.status != "skip":
                written_refs.add(w.ref)
            plan.writes.append(w)

    # 書かないセルの警告は、列と理由ごとにまとめる
    groups: dict[tuple[str, str], list[str]] = {}
    for w in plan.writes:
        if w.status == "skip":
            groups.setdefault((w.col, w.reason.split("（")[0] if "許可値にない値" not in w.reason else w.reason), []).append(w.ref)
    for (col, reason), refs in groups.items():
        where = refs[0] if len(refs) == 1 else f"{refs[0]}〜{refs[-1]}（{len(refs)}セル）"
        plan.warnings.append(f"{where}：{reason}のため書き込みません")

    if clear_sample and end:
        for r in range(start, end + 1):
            for col in profile.get("columns", {}):
                ref = f"{col}{r}"
                if ref in written_refs or ref in formulas or ref not in current:
                    continue
                if protected and ref not in unlocked:
                    continue
                if inspect.merged_anchor(analysis, ref) not in (None, ref):
                    continue
                plan.writes.append(CellWrite(ref, r, col, None, profile["columns"][col], "clear",
                                             "テンプレートに残っていた値を消去", current[ref]))
    plan.writes.sort(key=lambda w: (w.row, column_index_from_string(w.col)))
    return plan


def describe_sources(sources: dict[str, str | None], headers: dict[str, str]) -> list[str]:
    out = []
    for col, src in sources.items():
        if src and src.startswith(CONST_PREFIX):
            label = f"固定値「{src[len(CONST_PREFIX):]}」"
        elif src == TYPE_SOURCE:
            label = "WBSの種類"
        else:
            label = "連番" if src == "__seq__" else (MEANINGS.get(src, src) if src else "（書き込まない）")
        out.append(f"{col}列「{headers.get(col, '')}」← {label}")
    return out


def col_letter(i: int) -> str:
    return get_column_letter(i)
