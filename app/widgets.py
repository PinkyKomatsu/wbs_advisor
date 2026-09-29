"""画面で共有する部品：アプリの状態・QThread ワーカー・例外表示・シートのプレビューグリッド・案件属性の入力欄。"""
from __future__ import annotations

import functools
import logging
import traceback

from openpyxl.utils import get_column_letter
from PySide6.QtCore import QObject, QThread, Qt, Signal
from PySide6.QtGui import QBrush, QColor, QFont
from PySide6.QtWidgets import (QAbstractItemView, QCheckBox, QComboBox, QFormLayout, QGridLayout, QLineEdit,
                               QMessageBox, QSpinBox, QTableWidget, QTableWidgetItem, QWidget)

from core import config as config_mod
from core.db import Database
from excel_io import reader

log = logging.getLogger("wbsadvisor")

COLOR_HEADER = "#dbe8ff"
COLOR_DATA = "#fff8d6"
COLOR_SELECTED_COL = "#e6f4ea"
COLOR_TARGET = "#ffd8a8"


def error_message(e: BaseException) -> str:
    """利用者向けの日本語メッセージ。"""
    if isinstance(e, PermissionError):
        return "ファイルに書き込めません。Excel で開いていないか、保存先に書き込めるか確認してください。"
    if isinstance(e, FileNotFoundError):
        return f"ファイルが見つかりません：{getattr(e, 'filename', '') or e}"
    if e.__class__.__module__.startswith("pywintypes") or "com_error" in e.__class__.__name__:
        return "Excel の操作中にエラーが発生しました。Excel が応答しているか確認し、もう一度実行してください。"
    msg = str(e)
    return msg if msg and any("぀" <= ch <= "鿿" for ch in msg) else f"処理中にエラーが発生しました（{e.__class__.__name__}）。"


def show_error(parent, title: str, e: BaseException) -> None:
    # ログには例外の種類と場所だけを書き、メッセージ（セルの内容を含み得る）は書かない
    tb = traceback.extract_tb(e.__traceback__)
    where = " <- ".join(f"{f.filename.rsplit(chr(92), 1)[-1]}:{f.lineno}" for f in reversed(tb[-4:]))
    log.error("%s: %s at %s", title, e.__class__.__name__, where)
    QMessageBox.critical(parent, title, error_message(e))


def guarded(title: str):
    """スロットの例外を捕捉して日本語のダイアログを出すデコレーター。"""
    def deco(fn):
        @functools.wraps(fn)
        def wrapper(self, *args, **kwargs):
            try:
                return fn(self, *args, **kwargs)
            except Exception as e:  # すべて捕捉する
                show_error(self, title, e)
                return None
        return wrapper
    return deco


class AppState(QObject):
    formatsChanged = Signal()
    casesChanged = Signal()
    judgmentsChanged = Signal()
    settingsChanged = Signal()
    typeChanged = Signal()

    def __init__(self, db: Database | None = None):
        super().__init__()
        self.db = db or Database()
        self.cfg = config_mod.load()
        self.templates = config_mod.load_text_templates()
        self.judgments: list = []
        self.untranslated: dict = {}
        last = self.cfg.get("last_wbs_type")
        self.db.wbs_type = last if last in self.types() else self.types()[0]

    def types(self) -> list[str]:
        """WBS の種類（設計・構築・テスト など。設定で変更できる）。"""
        return [t for t in self.cfg.get("wbs_types") or [] if t] or ["設計", "構築", "テスト"]

    @property
    def wbs_type(self) -> str:
        return self.db.wbs_type

    def set_type(self, wbs_type: str):
        """WBS の種類を切り替える。ひな形・様式・過去事例・ルール・上書きはすべて種類ごと。"""
        if wbs_type == self.db.wbs_type:
            return
        self.db.wbs_type = wbs_type
        self.judgments, self.untranslated = [], {}
        self.cfg["last_wbs_type"] = wbs_type
        config_mod.save(self.cfg)
        log.info("WBS の種類を切り替えました")
        self.typeChanged.emit()
        self.formatsChanged.emit()
        self.casesChanged.emit()
        self.judgmentsChanged.emit()

    def save_config(self):
        config_mod.save(self.cfg)
        self.settingsChanged.emit()


class Worker(QThread):
    """重い処理をバックグラウンドで実行する。fn(progress) の progress(割合 0-100, 説明) で進捗を知らせる。"""
    progressed = Signal(int, str)
    succeeded = Signal(object)
    failed = Signal(object)

    def __init__(self, fn, parent=None):
        super().__init__(parent)
        self.fn = fn

    def run(self):
        try:
            result = self.fn(lambda pct, text="": self.progressed.emit(int(pct), text))
        except Exception as e:
            self.failed.emit(e)
        else:
            self.succeeded.emit(result)


def run_in_thread(owner, fn, on_done, title: str):
    """Worker を起動し、メインウィンドウの進捗バーに進捗を出す。"""
    win = owner.window()
    w = Worker(fn, owner)
    owner._workers = getattr(owner, "_workers", set())
    owner._workers.add(w)
    if hasattr(win, "begin_busy"):
        win.begin_busy(title)
        w.progressed.connect(win.set_progress)

    def done(result):
        owner._workers.discard(w)
        if hasattr(win, "end_busy"):
            win.end_busy()
        try:
            on_done(result)
        except Exception as e:
            show_error(owner, title, e)

    def fail(e):
        owner._workers.discard(w)
        if hasattr(win, "end_busy"):
            win.end_busy()
        show_error(owner, title, e)

    w.succeeded.connect(done)
    w.failed.connect(fail)
    w.start()
    return w


class SheetGrid(QTableWidget):
    """シートのプレビュー（列幅・結合を反映）。クリックしたセルを cellPicked(行, 列文字) で知らせる。"""
    cellPicked = Signal(int, str)

    def __init__(self, parent=None, max_rows: int = 200, max_cols: int = 30):
        super().__init__(parent)
        self.max_rows, self.max_cols = max_rows, max_cols
        self.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.setSelectionMode(QAbstractItemView.SingleSelection)
        self.cellClicked.connect(lambda r, c: self.cellPicked.emit(r + 1, get_column_letter(c + 1)))
        self.base_fill: dict[tuple[int, int], str | None] = {}

    def load(self, path, sheet: str):
        wb = reader.load(path, data_only=True)
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        n_rows, n_cols = min(ws.max_row, self.max_rows), min(ws.max_column, self.max_cols)
        self.clear()
        self.setRowCount(n_rows)
        self.setColumnCount(n_cols)
        self.setHorizontalHeaderLabels([get_column_letter(c) for c in range(1, n_cols + 1)])
        self.base_fill.clear()
        for r in range(1, n_rows + 1):
            for c in range(1, n_cols + 1):
                cell = ws.cell(r, c)
                item = QTableWidgetItem(reader.cell_text(cell.value))
                fill = None
                try:
                    if cell.fill is not None and cell.fill.fill_type == "solid" and isinstance(cell.fill.fgColor.rgb, str):
                        fill = "#" + cell.fill.fgColor.rgb[-6:]
                except (AttributeError, TypeError):
                    pass
                if fill:
                    item.setBackground(QBrush(QColor(fill)))
                if cell.font is not None and cell.font.b:
                    f = QFont(self.font())
                    f.setBold(True)
                    item.setFont(f)
                self.base_fill[(r, c)] = fill
                self.setItem(r - 1, c - 1, item)
        for m in ws.merged_cells.ranges:
            if m.min_row <= n_rows and m.min_col <= n_cols:
                self.setSpan(m.min_row - 1, m.min_col - 1, min(m.max_row, n_rows) - m.min_row + 1,
                             min(m.max_col, n_cols) - m.min_col + 1)
        for c in range(1, n_cols + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            w = dim.width if dim is not None and dim.width else 8.43
            self.setColumnWidth(c - 1, int(w * 7 + 5))
        wb.close()

    def highlight(self, profile: dict | None, targets: set[str] | None = None):
        """ヘッダー行（青）・データ範囲（黄）・書き込み先のセル（オレンジ）を色分けする。"""
        for (r, c), fill in self.base_fill.items():
            item = self.item(r - 1, c - 1)
            if item is not None:
                item.setBackground(QBrush(QColor(fill)) if fill else QBrush())
        if not profile:
            return
        hr, ds, de = profile.get("header_row"), profile.get("data_start"), profile.get("data_end")
        mapped = {self._col_index(col) for col in profile.get("columns", {})}
        for (r, c) in self.base_fill:
            item = self.item(r - 1, c - 1)
            if item is None:
                continue
            if hr and r == hr and c in mapped:
                item.setBackground(QBrush(QColor(COLOR_HEADER)))
            elif ds and r >= ds and (not de or r <= de) and c in mapped and r - ds < 60:
                item.setBackground(QBrush(QColor(COLOR_DATA)))
        for ref in targets or ():
            import re
            m = re.match(r"([A-Z]+)(\d+)", ref)
            r, c = int(m.group(2)), self._col_index(m.group(1))
            item = self.item(r - 1, c - 1)
            if item is not None:
                item.setBackground(QBrush(QColor(COLOR_TARGET)))

    @staticmethod
    def _col_index(col: str) -> int:
        from openpyxl.utils import column_index_from_string
        return column_index_from_string(col)


class AttrsEditor(QWidget):
    """案件属性（更新種別・病床規模・ベンダー・連携システム・年度・タグ）の入力欄。すべて任意。"""

    def __init__(self, cfg: dict, parent=None):
        super().__init__(parent)
        a = cfg.get("attributes", {})
        form = QFormLayout(self)
        form.setContentsMargins(0, 0, 0, 0)
        self.update_type = QComboBox()
        self.update_type.addItem("（未入力）", "")
        for t in a.get("update_types", []):
            self.update_type.addItem(t, t)
        self.beds = QSpinBox()
        self.beds.setRange(0, 3000)
        self.beds.setSpecialValueText("（未入力）")
        self.beds.setSuffix(" 床")
        self.vendor = QLineEdit()
        self.year = QSpinBox()
        self.year.setRange(0, 2100)
        self.year.setSpecialValueText("（未入力）")
        self.tags = QLineEdit()
        self.tags.setPlaceholderText("カンマ区切り")
        box = QWidget()
        grid = QGridLayout(box)
        grid.setContentsMargins(0, 0, 0, 0)
        self.systems: dict[str, QCheckBox] = {}
        for i, s in enumerate(a.get("systems", [])):
            cb = QCheckBox(s)
            self.systems[s] = cb
            grid.addWidget(cb, i // 4, i % 4)
        form.addRow("更新種別", self.update_type)
        form.addRow("病床規模", self.beds)
        form.addRow("ベンダー名", self.vendor)
        form.addRow("連携システム", box)
        form.addRow("実施年度", self.year)
        form.addRow("自由タグ", self.tags)

    def set_attrs(self, attrs: dict):
        attrs = attrs or {}
        i = self.update_type.findData(attrs.get("update_type", ""))
        self.update_type.setCurrentIndex(max(0, i))
        self.beds.setValue(int(attrs.get("beds") or 0))
        self.vendor.setText(attrs.get("vendor") or "")
        systems = set(attrs.get("systems") or [])
        for s, cb in self.systems.items():
            cb.setChecked(s in systems)
        self.year.setValue(int(attrs.get("year") or 0))
        self.tags.setText(", ".join(attrs.get("tags") or []))

    def attrs(self) -> dict:
        out = {}
        if self.update_type.currentData():
            out["update_type"] = self.update_type.currentData()
        if self.beds.value():
            out["beds"] = self.beds.value()
        if self.vendor.text().strip():
            out["vendor"] = self.vendor.text().strip()
        systems = [s for s, cb in self.systems.items() if cb.isChecked()]
        if systems:
            out["systems"] = systems
        if self.year.value():
            out["year"] = self.year.value()
        tags = [t.strip() for t in self.tags.text().replace("、", ",").split(",") if t.strip()]
        if tags:
            out["tags"] = tags
        return out


def ro_item(text, data=None) -> QTableWidgetItem:
    item = QTableWidgetItem("" if text is None else str(text))
    item.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
    if data is not None:
        item.setData(Qt.UserRole, data)
    return item
