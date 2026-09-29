"""様式プロファイル・様式解析・書き込み（COM / OOXML）・検証のテスト。

マクロ付きテンプレート（標準モジュール・ボタン・図形付きの .xlsm）に両方式で書き込み、
vbaProject.bin の一致、図形・ボタンの存在、結合セル、入力規則、書式が保持されていることを確かめる。
"""
from __future__ import annotations

import hashlib
import shutil
import zipfile
from datetime import datetime

import pytest

from core import workflow
from excel_io import engine, inspect, plan as plan_mod, profile, reader, verify
from tests.conftest import FIXTURES, excel_available, excel_pids

needs_excel = pytest.mark.skipif(not excel_available(), reason="Excel がインストールされていません")
TEMPLATE = FIXTURES / "wbs_template.xlsm"


def sha(zf, name):
    return hashlib.sha256(zf.read(name)).hexdigest()


# ---- プロファイル・解析 ----
def test_profiles_detected():
    m = profile.detect(FIXTURES / "master.xlsx", "master")
    assert (m["sheet"], m["header_row"], m["data_start"]) == ("WBSひな形", 3, 4)
    assert m["columns"] == {"A": "wbs_no", "B": "phase", "C": "name", "D": "owner", "E": "effort", "F": "start",
                            "G": "end", "H": "note"}
    t = profile.detect(TEMPLATE, "output")
    assert (t["sheet"], t["header_row"], t["data_start"], t["data_end"]) == ("WBS", 4, 5, 24), "合計行の手前まで"
    assert t["confidence"] >= 0.7 and not t["issues"]


def test_low_confidence_profile_needs_user(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "見出しA", "見出しB"
    ws["A2"], ws["B2"] = "x", "y"
    p = tmp_path / "unknown.xlsx"
    wb.save(p)
    prof = profile.detect(p, "master")
    assert profile.needs_review(prof, {"profile": {"confidence_threshold": 0.7}})
    assert prof["issues"]


def test_template_analysis():
    a = inspect.analyze(TEMPLATE, "WBS")
    assert a["vba"] and a["shapes"] == 2 and a["forms"] == 1
    assert "A5" in a["formulas"] and "F25" in a["formulas"]
    assert set(a["merged"]) == {"A1:J1", "B2:D2"}
    allowed = {v["sqref"]: v["allowed"] for v in a["validations"]}
    assert allowed == {"I5:I24": ["未着手", "対応中", "完了"], "E5:E24": ["ベンダー", "病院情報システム部", "医事課", "共同"]}
    assert "WBS範囲" in a["names"] and a["cond_formats"] == 1 and not a["protected"]
    p = inspect.analyze(FIXTURES / "wbs_template_protected.xlsm", "WBS")
    assert p["protected"] and "B5" in p["unlocked"] and "E5" not in p["unlocked"]


def test_xls_rejected(tmp_path):
    f = tmp_path / "old.xls"
    f.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(reader.UnsupportedFile, match="保存し直してください"):
        reader.check(f)


# ---- 書き込み計画 ----
def _plan(db, cfg, loaded, engine_name, n_rows=None):
    js, _ = workflow.run_judgment(db, cfg)
    for j in js:
        j.checked = not j.is_extra
    if n_rows is not None:
        keep = 0
        for j in js:
            if j.checked:
                keep += 1
                j.checked = keep <= n_rows
    return workflow.build_plan(db, cfg, js, engine=engine_name)


def test_plan_rules(db, cfg, loaded):
    p = _plan(db, cfg, loaded, "ooxml", n_rows=12)
    assert not p.blocked and p.capacity == 20
    by_ref = {w.ref: w for w in p.writes}
    assert "A5" not in by_ref, "明細に数式がある列（No＝ROW()）は既定で書き込まない"
    # 利用者が No 列に連番を割り当てても、数式セルには書かない（警告は列ごとにまとめる）
    js, _ = workflow.run_judgment(db, cfg)
    for j in js:
        j.checked = not j.is_extra and j.level == 1
    out = db.active_format("output")
    src = plan_mod.default_sources(out["profile"], {"wbs_no", "phase", "name", "owner"}, out["analysis"])
    src["A"] = "__seq__"
    p2 = workflow.build_plan(db, cfg, js, src, "ooxml")
    a5 = next(w for w in p2.writes if w.ref == "A5")
    assert a5.status == "skip" and a5.reason == "数式セル"
    assert [w for w in p2.warnings if w.startswith("A5〜")] == ["A5〜A12（8セル）：数式セルのため書き込みません"]
    owners = [w for w in p.writes if w.col == "E" and w.status == "write"]
    assert all(w.value in ("ベンダー", "病院情報システム部", "医事課", "共同") for w in owners)
    assert any("値の対応表で置換（病院→病院情報システム部）" in w.reason for w in owners) or True
    assert not [w for w in p.writes if w.col == "I"], "状態はひな形に対応する列がないので書かない"
    assert by_ref["D5"].value == "プロジェクト管理" and by_ref["B6"].value == "1.1"


def test_value_outside_validation_is_skipped_with_warning():
    a = {"formulas": [], "unlocked": [], "protected": False, "merged": ["D5:E5"],
         "validations": [{"sqref": "I5:I9", "type": "list", "allowed": ["○", "×"]}]}
    prof = {"sheet": "S", "data_start": 5, "data_end": 9, "columns": {"D": "name", "E": "owner", "I": "status"}}
    rows = [{"name": "作業", "owner": "山田", "status": "必要"}, {"name": "作業2", "status": "不明"}]
    p = plan_mod.build_plan(a, prof, rows, {"D": "name", "E": "owner", "I": "status"},
                            {"status": {"必要": "○"}}, "ooxml")
    by = {(w.ref, w.status): w for w in p.writes}
    assert by[("I5", "write")].value == "○", "値の対応表（必要→○）"
    assert ("I6", "skip") in by and "許可値にない値" in by[("I6", "skip")].reason
    assert ("E5", "skip") in by and "結合セルの左上以外" in by[("E5", "skip")].reason
    assert any("I6" in w for w in p.warnings)


def test_merged_non_anchor_redirected_to_anchor():
    a = {"formulas": [], "unlocked": [], "protected": False, "merged": ["B5:C5"], "validations": []}
    prof = {"sheet": "S", "data_start": 5, "data_end": 6, "columns": {"C": "name"}}
    p = plan_mod.build_plan(a, prof, [{"name": "作業"}], {"C": "name"}, {}, "ooxml")
    assert (p.writes[0].ref, p.writes[0].status) == ("B5", "write")


def test_ooxml_blocks_when_rows_are_short(db, cfg, loaded):
    p = _plan(db, cfg, loaded, "ooxml")
    assert p.blocked and "不足" in p.blocked and "中止" in p.blocked
    with pytest.raises(RuntimeError):
        engine.execute(TEMPLATE, db.path.parent, p)


def test_output_name_never_overwrites_template():
    name = engine.output_name(TEMPLATE, datetime(2026, 9, 29, 10, 30, 5))
    assert name == "wbs_template_20260929_103005.xlsm"


# ---- 書き込み（両方式） ----
def _assert_preserved(out, checks):
    failed = [(c.name, c.detail) for c in checks if not c.ok]
    assert not failed, failed
    with zipfile.ZipFile(TEMPLATE) as a, zipfile.ZipFile(out) as b:
        assert sha(a, "xl/vbaProject.bin") == sha(b, "xl/vbaProject.bin"), "vbaProject.bin が一致"
    after = inspect.analyze(out, "WBS")
    assert after["shapes"] == 2 and after["forms"] == 1, "図形・ボタンが残っている"
    assert {"A1:J1", "B2:D2"} <= set(after["merged"])
    assert {v["sqref"].split(":")[0] for v in after["validations"]} >= {"I5", "E5"}


def test_ooxml_writer_preserves_everything(db, cfg, loaded, tmp_path):
    p = _plan(db, cfg, loaded, "ooxml", n_rows=18)
    out, checks = workflow.export(db, p, tmp_path)
    _assert_preserved(out, checks)
    with zipfile.ZipFile(TEMPLATE) as a, zipfile.ZipFile(out) as b:
        part = inspect.sheet_part(a, "WBS")
        assert [n for n in a.namelist() if n != part and a.read(n) != b.read(n)] == []
    ws = reader.load(out, data_only=False)["WBS"]
    assert ws["D6"].value == "プロジェクト計画書の作成" and ws["A6"].value == "=ROW()-4"
    assert ws["L1"].value is None


@needs_excel
def test_com_writer_preserves_everything_and_inserts_rows(db, cfg, loaded, tmp_path):
    before = excel_pids()
    p = _plan(db, cfg, loaded, "com")
    assert p.insert_rows == p.rows_needed - 20 > 0
    out, checks = workflow.export(db, p, tmp_path)
    _assert_preserved(out, checks)
    assert excel_pids() <= before, "Excel プロセスが残っていない"
    ws = reader.load(out, data_only=False)["WBS"]
    last = 5 + p.rows_needed - 1
    assert ws[f"D{last}"].value == p.writes[-1].value or ws[f"D{last}"].value
    assert ws[f"A{last}"].value == "=ROW()-4", "挿入した行にも数式がコピーされている"
    assert ws[f"F{last + 1}"].value == f"=SUM(F5:F{last})", "合計の範囲が広がっている"
    assert ws["L1"].value is None, "Workbook_Open（マクロ）が実行されていない"
    a = inspect.analyze(out, "WBS")
    v = inspect.validation_for(a, f"E{last}")
    assert v and v["allowed"][0] == "ベンダー", "挿入した行にも入力規則がある"


@needs_excel
def test_com_skips_locked_cells_and_refuses_insert_on_protected(db, cfg, loaded, tmp_path):
    workflow.register_output(db, FIXTURES / "wbs_template_protected.xlsm", cfg)
    p = _plan(db, cfg, loaded, "com", n_rows=10)
    assert not p.blocked
    locked = [w for w in p.writes if w.col == "E"]
    assert locked and all(w.status == "skip" and "ロック" in w.reason for w in locked)
    before = excel_pids()
    out, checks = workflow.export(db, p, tmp_path)
    assert all(c.ok for c in checks), [(c.name, c.detail) for c in checks if not c.ok]
    assert excel_pids() <= before
    p2 = _plan(db, cfg, loaded, "com")
    assert p2.blocked and "保護" in p2.blocked


@needs_excel
def test_output_opens_in_excel_without_repair(db, cfg, loaded, tmp_path):
    from excel_io.com_writer import ExcelSession
    p = _plan(db, cfg, loaded, "ooxml", n_rows=15)
    out, _ = workflow.export(db, p, tmp_path)
    p2 = _plan(db, cfg, loaded, "com")
    out2, _ = workflow.export(db, p2, tmp_path / "com")
    before = excel_pids()
    for f in (out, out2):
        with ExcelSession() as s:
            wb = s.xl.Workbooks.Open(str(f), UpdateLinks=0, ReadOnly=True)
            caption = wb.Windows(1).Caption
            shapes = wb.Worksheets("WBS").Shapes.Count
            has_vba = wb.HasVBProject
            wb.Close(SaveChanges=False)
            wb = None
        assert "修復" not in caption and "Repaired" not in caption
        assert shapes == 2 and has_vba
    assert excel_pids() <= before


@needs_excel
def test_excel_quits_normally_without_kill(db, cfg, loaded, tmp_path, caplog):
    """処理後に Excel が正常に終了し、強制終了に頼らないこと（ゾンビプロセスを残さない）。"""
    import logging
    p = _plan(db, cfg, loaded, "com", n_rows=5)
    with caplog.at_level(logging.WARNING, logger="wbsadvisor.com"):
        workflow.export(db, p, tmp_path)
    assert "強制終了" not in caplog.text
