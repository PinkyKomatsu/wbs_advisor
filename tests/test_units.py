"""WBS の種類ごとの蓄積と、「大分類・中分類・作業項目」→「管理単位（中２）・（中３）・（小）」の転記のテスト。"""
from __future__ import annotations

import sqlite3

import pytest

from core import workflow
from core.db import Database
from excel_io import plan as plan_mod, reader
from tests.conftest import FIXTURES


@pytest.fixture()
def units(db, cfg):
    """「設計」WBS として、ひな形・実用WBS様式・過去事例 2 件を登録した状態。"""
    db.wbs_type = "設計"
    workflow.register_master(db, FIXTURES / "master_units.xlsx", cfg)
    workflow.register_output(db, FIXTURES / "template_units.xlsx", cfg)
    for name in ("units_A", "units_B"):
        workflow.import_case(db, FIXTURES / f"{name}.xlsx", cfg, name=name)
    return db


# ---- WBS の種類 ----
def test_cases_are_kept_per_wbs_type(units, cfg):
    db = units
    assert [c.name for c in db.cases()] == ["units_A", "units_B"]
    db.wbs_type = "テスト"
    assert db.cases() == [] and db.active_format("master") is None and db.active_format("output") is None
    assert workflow.master_items(db) == []
    # 同じファイルでも、別の種類なら取り込める
    workflow.register_master(db, FIXTURES / "master_units.xlsx", cfg)
    workflow.import_case(db, FIXTURES / "units_A.xlsx", cfg, name="テストの事例")
    assert [c.name for c in db.cases()] == ["テストの事例"]
    db.wbs_type = "設計"
    assert [c.name for c in db.cases()] == ["units_A", "units_B"], "他の種類の事例は混ざらない"
    counts = db.counts_by_type()
    assert counts["設計"]["cases"] == 2 and counts["テスト"]["cases"] == 1


def test_duplicate_detected_per_type_and_can_be_forced(units, cfg):
    with pytest.raises(workflow.DuplicateCase):
        workflow.import_case(units, FIXTURES / "units_A.xlsx", cfg)
    workflow.import_case(units, FIXTURES / "units_A.xlsx", cfg, allow_duplicate=True, name="二重")
    assert len(units.cases()) == 3


def test_rules_synonyms_overrides_are_per_type(units, cfg):
    from core.models import Rule
    db = units
    key = workflow.master_items(db)[0].key
    db.save_rule(Rule(None, "設計のルール", targets=[key]))
    db.add_synonym("ほげ", key)
    db.add_override(key, "不要", "必要")
    db.wbs_type = "構築"
    assert db.rules() == [] and db.synonyms() == {} and db.overrides() == {}
    db.wbs_type = "設計"
    assert len(db.rules()) == 1 and db.synonyms() and db.overrides() == {key: "不要"}


# ---- 大分類・中分類・作業項目 ----
def test_master_items_have_categories_filled_down(units):
    items = workflow.master_items(units)
    assert [i.name for i in items] == [n for _, _, n, _ in __import__("tests.fixtures.make_fixtures", fromlist=["x"]).UNIT_MASTER]
    pacs = next(i for i in items if i.name == "PACS連携仕様の作成")
    assert (pacs.values["cat1"], pacs.values["cat2"]) == ("詳細設計", "連携設計"), "空欄の分類は上の行から引き継ぐ"
    assert pacs.key == "詳細設計連携設計/pacs連携仕様の作成", "識別子は 大分類＋中分類／作業項目（正規化）"


def test_case_rows_read_from_management_units(units):
    rows = [r for r in units.case_rows() if r.name == "PACS連携仕様の作成"]
    assert len(rows) == 1 and rows[0].values["cat1"] == "詳細設計" and rows[0].values["cat2"] == "連携設計"
    assert rows[0].status == "auto"
    variant = next(r for r in units.case_rows() if r.name == "画面一覧作成")
    assert variant.status == "auto" and variant.item_key == next(
        i.key for i in workflow.master_items(units) if i.name == "画面一覧の作成")


def test_judgment_and_copy_with_categories(units, cfg):
    js, _ = workflow.run_judgment(units, cfg)
    by = {j.name: j for j in js}
    assert by["PACS連携仕様の作成"].rate == pytest.approx(0.5) and by["画面一覧の作成"].rate == 1.0
    assert (by["点数マスタ移行設計"].cat1, by["点数マスタ移行設計"].cat2) == ("詳細設計", "マスタ設計")
    from core import config, texts
    tsv = texts.copy_text([by["点数マスタ移行設計"]], "ja", "tsv", config.load_text_templates())
    assert tsv.split("\n")[0].split("\t")[:3] == ["大分類", "中分類", "作業項目"], "WBS番号のない WBS では列を省く"
    assert tsv.split("\n")[1].split("\t")[:3] == ["詳細設計", "マスタ設計", "点数マスタ移行設計"]
    en = texts.copy_text([by["点数マスタ移行設計"]], "en", "tsv", config.load_text_templates())
    assert "【未翻訳: 詳細設計】" in en, "分類名も用語辞書で英訳（未登録なら未翻訳）"


def _plan(db, cfg, sources=None):
    js, _ = workflow.run_judgment(db, cfg)
    for j in js:
        j.checked = True
    return workflow.build_plan(db, cfg, js, sources, "ooxml")


def test_transfer_to_management_unit_columns(units, cfg, tmp_path):
    p = _plan(units, cfg)
    assert not p.blocked
    by = {w.ref: w.value for w in p.writes if w.status == "write"}
    # 大分類→管理単位（中２）＝C、中分類→（中３）＝D、作業項目→（小）＝E
    assert (by["C4"], by["D4"], by["E4"]) == ("基本設計", "画面設計", "画面一覧の作成")
    assert (by["C9"], by["D9"], by["E9"]) == ("詳細設計", "連携設計", "PACS連携仕様の作成")
    # 管理単位（大）・（中１）は、過去事例で最も多く使われた値を全行に書く
    assert (by["A4"], by["B4"], by["A10"], by["B10"]) == ("医事会計更新", "設計", "医事会計更新", "設計")
    out, checks = workflow.export(units, p, tmp_path)
    assert all(c.ok for c in checks), [(c.name, c.detail) for c in checks if not c.ok]
    ws = reader.load(out)["WBS"]
    assert [ws[f"E{r}"].value for r in range(4, 11)] == [
        "画面一覧の作成", "画面レイアウトの作成", "帳票一覧の作成", "帳票レイアウトの作成",
        "電子カルテ連携仕様の作成", "PACS連携仕様の作成", "点数マスタ移行設計"]
    assert ws["C10"].value == "詳細設計" and ws["D10"].value == "マスタ設計"


def test_unit_candidates_from_past_cases(units):
    cands = workflow.unit_candidates(units, "unit_l")
    assert cands[0][0] == "医事会計更新" and cands[0][1] > 0
    assert any(v == "設計" for v, _ in cands), "WBS の種類名も候補に入る"
    m1 = workflow.unit_candidates(units, "unit_m1")
    assert m1[0][0] == "設計"
    out = units.active_format("output")
    src = workflow.default_sources(units, out["profile"], {"cat1", "cat2", "name"}, out["analysis"])
    assert src["A"] == plan_mod.CONST_PREFIX + "医事会計更新" and src["B"] == plan_mod.CONST_PREFIX + "設計"
    units.wbs_type = "構築"   # 事例のない種類では、既定では書かない（候補は種類名のみ）
    assert workflow.unit_candidates(units, "unit_l") == [("構築", 0)]
    units.wbs_type = "設計"


def test_fixed_value_and_wbs_type_sources(units, cfg):
    out = units.active_format("output")
    src = plan_mod.default_sources(out["profile"], {"cat1", "cat2", "name", "owner"}, out["analysis"])
    src["A"] = plan_mod.CONST_PREFIX + "医事会計システム更新"
    src["B"] = plan_mod.TYPE_SOURCE
    p = _plan(units, cfg, src)
    by = {w.ref: w.value for w in p.writes if w.status == "write"}
    assert by["A4"] == "医事会計システム更新" and by["B4"] == "設計" and by["B10"] == "設計"


def test_categories_can_be_written_only_when_changed(units, cfg):
    cfg["output"]["repeat_categories"] = False
    p = _plan(units, cfg)
    vals = {w.ref: w.value for w in p.writes if w.status == "write"}
    assert vals["C4"] == "基本設計" and "C5" not in vals, "同じ大分類は 2 行目から空欄"
    assert vals["D6"] == "帳票設計" and "D7" not in vals
    assert vals["C8"] == "詳細設計"


# ---- 書式の検証（COM 方式） ----
def test_same_style_ignores_representation_but_catches_real_changes():
    import openpyxl
    from openpyxl.styles import Border, Font, PatternFill, Side
    from excel_io.verify import same_style
    ws = openpyxl.Workbook().active
    a, b = ws["A1"], ws["B1"]
    a.font = Font(name="Calibri", sz=11, scheme="minor")
    b.font = Font(name="ＭＳ Ｐゴシック", sz=11)                     # テーマのフォントを Excel が名前で書き出した
    a.fill = PatternFill("solid", fgColor="00D9E1F2")
    b.fill = PatternFill("solid", fgColor="FFD9E1F2")                 # 透明度の表記だけ違う
    assert same_style(a, b)
    for change in (lambda c: setattr(c, "font", Font(name="ＭＳ Ｐゴシック", sz=11, b=True)),
                   lambda c: setattr(c, "border", Border(left=Side(style="thin"))),
                   lambda c: setattr(c, "fill", PatternFill("solid", fgColor="FFFF0000")),
                   lambda c: setattr(c, "number_format", "0.0")):
        c = ws["C1"]
        c.font, c.fill = Font(name="ＭＳ Ｐゴシック", sz=11), PatternFill("solid", fgColor="FFD9E1F2")
        c.border, c.number_format = Border(), "General"
        change(c)
        assert not same_style(a, c)
    d, e = ws["D1"], ws["E1"]
    d.font, e.font = Font(name="Meiryo", sz=10), Font(name="ＭＳ Ｐゴシック", sz=10)   # テーマ以外のフォントの変更は差
    assert not same_style(d, e)


@pytest.mark.skipif(not __import__("tests.conftest", fromlist=["x"]).excel_available(), reason="Excel がありません")
def test_com_on_openpyxl_made_template(units, cfg, tmp_path):
    js, _ = workflow.run_judgment(units, cfg)
    for j in js:
        j.checked = True
    p = workflow.build_plan(units, cfg, js, None, "com")
    out, checks = workflow.export(units, p, tmp_path)
    assert all(c.ok for c in checks), [(c.name, c.detail) for c in checks if not c.ok]
    ws = reader.load(out)["WBS"]
    assert ws["E9"].value == "PACS連携仕様の作成" and ws["C9"].value == "詳細設計"


# ---- 古いデータベースの移行 ----
def test_old_database_is_migrated(tmp_path):
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.executescript("""
        CREATE TABLE formats(id INTEGER PRIMARY KEY, kind TEXT, name TEXT, file TEXT, source TEXT, sha256 TEXT,
                             profile TEXT, analysis TEXT, created TEXT, active INTEGER DEFAULT 1);
        CREATE TABLE cases(id INTEGER PRIMARY KEY, name TEXT NOT NULL, source TEXT, sha256 TEXT UNIQUE, imported TEXT, attrs TEXT);
        CREATE TABLE case_rows(id INTEGER PRIMARY KEY, case_id INTEGER NOT NULL REFERENCES cases(id) ON DELETE CASCADE,
                               seq INTEGER, row INTEGER, wbs_no TEXT, phase TEXT, name TEXT, norm TEXT, values_json TEXT,
                               item_key TEXT, method TEXT, score REAL, status TEXT, candidates TEXT, prev_key TEXT);
        CREATE TABLE synonyms(norm TEXT PRIMARY KEY, item_key TEXT NOT NULL, created TEXT);
        CREATE TABLE rules(id INTEGER PRIMARY KEY, name TEXT NOT NULL, name_en TEXT, enabled INTEGER, priority INTEGER,
                           conditions TEXT, decision TEXT, targets TEXT, keyword TEXT, created TEXT);
        CREATE TABLE overrides(id INTEGER PRIMARY KEY, item_key TEXT NOT NULL, decision TEXT, previous TEXT, note TEXT, at TEXT);
        INSERT INTO cases(id, name, sha256) VALUES(1, '旧事例', 'abc');
        INSERT INTO case_rows(case_id, seq, name) VALUES(1, 0, '作業');
        INSERT INTO synonyms VALUES('x', 'k', '');
        INSERT INTO rules(name) VALUES('旧ルール');
    """)
    con.commit()
    con.close()
    db = Database(path)
    assert db.wbs_type == "設計"
    fk_sql = db.conn.execute("SELECT sql FROM sqlite_master WHERE name = 'case_rows'").fetchone()[0]
    assert "REFERENCES cases(" in fk_sql and "cases_old" not in fk_sql and "cases_new" not in fk_sql
    assert [c.name for c in db.cases()] == ["旧事例"] and len(db.case_rows()) == 1
    assert db.synonyms() == {"x": "k"} and [r.name for r in db.rules()] == ["旧ルール"]
    db.wbs_type = "テスト"
    db.add_case("同じハッシュ", "", "abc", {}, [])     # UNIQUE をやめたので種類が違えば同じファイルを登録できる
    assert [c.name for c in db.cases()] == ["同じハッシュ"]
    db.close()
