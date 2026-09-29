"""テスト用のダミーファイルを作る（Excel が必要。開発 PC で 1 回だけ実行し、結果を tests/fixtures に保存する）。

    .venv\\Scripts\\python tests\\fixtures\\make_fixtures.py

作るもの:
  master.xlsx                 ひな形（医事会計システム更新の全作業項目、階層付き）
  case_A.xlsx / case_B.xlsx / case_C.xlsx
                              実用WBS（過去事例）。B は WBS 番号の振り直しと名称の表記ゆれ、C はひな形外の追加項目を含む
  wbs_template.xlsm           実用WBS様式（出力テンプレート）。標準モジュール・ボタン・図形・入力規則・数式・
                              条件付き書式・名前定義・印刷設定・Workbook_Open（開いたら L1 に書き込む）を含む
  wbs_template_protected.xlsm 上と同じでシート保護あり（B〜D・F〜J 列の明細だけロック解除）

マクロ付き .xlsm をプログラムで作るため、実行中だけ HKCU の
「VBA プロジェクト オブジェクト モデルへのアクセスを信頼する」（AccessVBOM）を 1 にし、終了時に元に戻す。
"""
from __future__ import annotations

import sys
import winreg
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent

HEADERS = ["No", "WBS番号", "工程", "作業項目", "担当", "予定工数", "開始予定", "終了予定", "状態", "備考"]
HEADER_ROW = 4
DATA_START = 5
DATA_END = 24          # テンプレートの明細は 20 行
OWNERS = ["ベンダー", "病院情報システム部", "医事課", "共同"]

# (WBS番号, 工程, 作業項目, 担当, 予定工数, 備考)
MASTER = [
    ("1", "プロジェクト管理", "プロジェクト管理", "", None, ""),
    ("1.1", "プロジェクト管理", "プロジェクト計画書の作成", "ベンダー", 5, ""),
    ("1.2", "プロジェクト管理", "キックオフ会議", "共同", 1, ""),
    ("1.3", "プロジェクト管理", "定例会議の運営", "共同", 10, "月2回"),
    ("1.4", "プロジェクト管理", "課題管理", "ベンダー", 4, ""),
    ("2", "要件定義", "要件定義", "", None, ""),
    ("2.1", "要件定義", "現行業務の調査", "ベンダー", 8, ""),
    ("2.2", "要件定義", "医事会計業務要件の定義", "共同", 10, ""),
    ("2.3", "要件定義", "レセプト要件の定義", "医事課", 6, ""),
    ("2.4", "要件定義", "帳票要件の定義", "医事課", 4, ""),
    ("2.5", "要件定義", "DPC要件の定義", "医事課", 4, "DPC対象病院のみ"),
    ("2.6", "要件定義", "要件定義書の承認", "病院", 1, ""),
    ("3", "設計", "設計", "", None, ""),
    ("3.1", "設計", "基本設計", "ベンダー", 15, ""),
    ("3.2", "設計", "電子カルテ連携設計", "ベンダー", 6, ""),
    ("3.3", "設計", "PACS連携設計", "ベンダー", 4, ""),
    ("3.4", "設計", "検査システム連携設計", "ベンダー", 4, ""),
    ("3.5", "設計", "点数マスタ移行設計", "ベンダー", 3, ""),
    ("4", "開発・テスト", "開発・テスト", "", None, ""),
    ("4.1", "開発・テスト", "カスタマイズ開発", "ベンダー", 30, ""),
    ("4.2", "開発・テスト", "単体テスト", "ベンダー", 10, ""),
    ("4.3", "開発・テスト", "結合テスト", "ベンダー", 10, ""),
    ("4.4", "開発・テスト", "電子カルテ連携テスト", "共同", 5, ""),
    ("4.5", "開発・テスト", "PACS連携テスト", "共同", 3, ""),
    ("4.6", "開発・テスト", "検査システム連携テスト", "共同", 3, ""),
    ("4.7", "開発・テスト", "総合テスト", "共同", 10, ""),
    ("4.8", "開発・テスト", "性能テスト", "ベンダー", 3, ""),
    ("5", "データ移行", "データ移行", "", None, ""),
    ("5.1", "データ移行", "移行計画の策定", "ベンダー", 3, ""),
    ("5.2", "データ移行", "患者基本情報の移行", "ベンダー", 5, ""),
    ("5.3", "データ移行", "未収金データの移行", "ベンダー", 3, ""),
    ("5.4", "データ移行", "移行リハーサル", "共同", 4, "2回実施"),
    ("6", "教育・運用準備", "教育・運用準備", "", None, ""),
    ("6.1", "教育・運用準備", "操作マニュアルの作成", "ベンダー", 5, ""),
    ("6.2", "教育・運用準備", "職員向け操作研修", "共同", 6, ""),
    ("6.3", "教育・運用準備", "運用手順書の作成", "医事課", 4, ""),
    ("7", "本番移行", "本番移行", "", None, ""),
    ("7.1", "本番移行", "稼働判定会議", "共同", 1, ""),
    ("7.2", "本番移行", "本番データ移行", "ベンダー", 2, ""),
    ("7.3", "本番移行", "切り戻し手順の確認", "ベンダー", 1, ""),
    ("8", "稼働後フォロー", "稼働後フォロー", "", None, ""),
    ("8.1", "稼働後フォロー", "稼働後の問い合わせ対応", "ベンダー", 10, ""),
    ("8.2", "稼働後フォロー", "初回レセプト請求の支援", "共同", 3, ""),
    ("8.3", "稼働後フォロー", "稼働後評価", "病院情報システム部", 2, ""),
]

# 過去事例：採用したひな形項目（WBS番号）と、表記ゆれ・追加項目
CASE_A = {  # リプレース・400床・ベンダーX・電子カルテ/PACS/検査
    "items": [n for n, *_ in MASTER],  # すべて採用
    "rename": {}, "renumber": False, "extra": [],
}
CASE_B = {  # バージョンアップ・200床・ベンダーY・電子カルテのみ
    "items": ["1", "1.1", "1.2", "1.3", "2", "2.1", "2.2", "2.3", "2.6", "3", "3.1", "3.2", "4", "4.1", "4.2",
              "4.3", "4.4", "4.7", "6", "6.1", "6.2", "7", "7.1", "7.2", "8", "8.1"],
    "rename": {"2.1": "現行業務調査", "6.2": "職員操作研修会"},   # 表記ゆれ（類似度 高 / 中）
    "renumber": True, "extra": [],
}
CASE_C = {  # リプレース・500床・ベンダーX・電子カルテ/PACS
    "items": ["1", "1.1", "1.2", "1.3", "1.4", "2", "2.1", "2.2", "2.3", "2.4", "2.6", "3", "3.1", "3.2", "3.3",
              "3.5", "4", "4.1", "4.2", "4.3", "4.4", "4.5", "4.7", "5", "5.1", "5.2", "5.4", "6", "6.1", "6.2",
              "7", "7.1", "7.2", "7.3", "8", "8.1", "8.2"],
    "rename": {}, "renumber": False,
    "extra": [("6.2", "院内説明会の実施", "病院情報システム部", 2)],   # 6.2 の後ろに追加
}

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")


def build_master(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WBSひな形"
    ws.merge_cells("A1:H1")
    ws["A1"] = "医事会計システム更新 WBSひな形"
    ws["A1"].font = Font(bold=True, size=14)
    heads = ["WBS番号", "工程", "作業項目", "担当", "予定工数(人日)", "開始日", "終了日", "備考"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(3, i, h)
        c.font, c.fill, c.border = Font(bold=True), HEAD_FILL, BOX
    for r, (no, phase, name, owner, effort, note) in enumerate(MASTER, 4):
        vals = [no, phase, name, owner, effort, None, None, note]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border = BOX
            if "." not in no and i <= 3:
                c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (9, 16, 30, 16, 12, 11, 11, 20)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def build_case(path: Path, spec: dict, title: str):
    by_no = {m[0]: m for m in MASTER}
    rows = []
    for no in spec["items"]:
        _, phase, name, owner, effort, note = by_no[no]
        rows.append([no, phase, spec["rename"].get(no, name), owner, effort, note])
        for after, xname, xowner, xeffort in spec["extra"]:
            if after == no:
                rows.append([no + ".1", phase, xname, xowner, xeffort, "追加"])
    if spec["renumber"]:   # 「2-1」形式に振り直す（番号では照合できない）
        major, minor = 0, 0
        for row in rows:
            if "." not in row[0]:
                major, minor = major + 1, 0
                row[0] = f"{major}"
            else:
                minor += 1
                row[0] = f"{major}-{minor}"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WBS"
    ws.merge_cells("A1:J1")
    ws["A1"] = "医事会計システム更新 WBS"
    ws["A2"], ws["B2"] = "案件名", title
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(HEADER_ROW, i, h)
        c.font, c.fill, c.border = Font(bold=True), HEAD_FILL, BOX
    for r, (no, phase, name, owner, effort, note) in enumerate(rows, DATA_START):
        vals = [r - HEADER_ROW, no, phase, name, owner if owner != "病院" else "病院情報システム部",
                effort, None, None, "完了", note]
        for i, v in enumerate(vals, 1):
            ws.cell(r, i, v).border = BOX
    wb.save(path)


# ---------------------------------------------------------------------------
# マクロ付きテンプレート（Excel COM で作る）
# ---------------------------------------------------------------------------

VBA_MODULE = '''Attribute VB_Name = "Module1"
Public Sub RenumberRows()
    Dim r As Long
    For r = 5 To 24
        Worksheets("WBS").Cells(r, 1).Formula = "=ROW()-4"
    Next r
End Sub
'''
VBA_WORKBOOK_OPEN = '''Private Sub Workbook_Open()
    Worksheets("WBS").Range("L1").Value = "マクロが実行されました"
End Sub
'''

def quit_and_wait(xl, timeout_ms: int = 15000) -> None:
    """Excel を終了させ、プロセスが完全に終わるまで待つ。
    （Excel 2013 は終了処理の最後にトラストセンターの設定をレジストリへ書き戻すため、待たずに戻すと上書きされる）"""
    import win32api
    import win32con
    import win32event
    import win32process
    try:
        _, pid = win32process.GetWindowThreadProcessId(xl.Hwnd)
    except Exception:
        pid = None
    xl.Quit()
    if pid:
        try:
            h = win32api.OpenProcess(win32con.SYNCHRONIZE | win32con.PROCESS_TERMINATE, False, pid)
        except Exception:
            return
        if win32event.WaitForSingleObject(h, timeout_ms) != win32event.WAIT_OBJECT_0:
            win32api.TerminateProcess(h, 1)
        win32api.CloseHandle(h)


def excel_version() -> str:
    """インストールされている Excel のバージョン（例 "15.0" = Excel 2013、"16.0" = 2016 以降）。"""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    try:
        return str(xl.Version)
    finally:
        quit_and_wait(xl)
        del xl
        pythoncom.CoUninitialize()


class TrustVBOM:
    """実行中だけ AccessVBOM=1 にして、終了時に元の値に戻す（キーは Excel のバージョンごと）。"""

    def __init__(self, version: str):
        self.path = rf"Software\Microsoft\Office\{version}\Excel\Security"

    def __enter__(self):
        self.key = winreg.CreateKeyEx(winreg.HKEY_CURRENT_USER, self.path, 0, winreg.KEY_READ | winreg.KEY_WRITE)
        try:
            self.old = winreg.QueryValueEx(self.key, "AccessVBOM")[0]
        except FileNotFoundError:
            self.old = None
        winreg.SetValueEx(self.key, "AccessVBOM", 0, winreg.REG_DWORD, 1)
        return self

    def __exit__(self, *exc):
        if self.old is None:
            try:
                winreg.DeleteValue(self.key, "AccessVBOM")
            except FileNotFoundError:
                pass
        else:
            winreg.SetValueEx(self.key, "AccessVBOM", 0, winreg.REG_DWORD, self.old)
        try:
            now = winreg.QueryValueEx(self.key, "AccessVBOM")[0]
        except FileNotFoundError:
            now = None
        winreg.CloseKey(self.key)
        if now != self.old:
            raise RuntimeError(f"AccessVBOM を元に戻せませんでした（現在 {now}、元 {self.old}）")


def build_template(path: Path, protected_path: Path):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Add()
        while wb.Worksheets.Count > 1:
            wb.Worksheets(wb.Worksheets.Count).Delete()
        ws = wb.Worksheets(1)
        ws.Name = "WBS"
        lists = wb.Worksheets.Add(After=ws)
        lists.Name = "リスト"
        for i, o in enumerate(OWNERS, 1):
            lists.Cells(i, 1).Value = o

        ws.Range("A1:J1").Merge()
        ws.Range("A1").Value = "医事会計システム更新 WBS"
        ws.Range("A1").Font.Bold = True
        ws.Range("A1").Font.Size = 14
        ws.Range("A2").Value = "案件名"
        ws.Range("B2:D2").Merge()
        ws.Range("B2:D2").Borders.LineStyle = 1
        for i, h in enumerate(HEADERS, 1):
            c = ws.Cells(HEADER_ROW, i)
            c.Value = h
            c.Font.Bold = True
            c.Interior.Color = 0xF2E1D9  # BGR
        rng = ws.Range(f"A{HEADER_ROW}:J{DATA_END}")
        rng.Borders.LineStyle = 1
        for r in range(DATA_START, DATA_END + 1):
            ws.Cells(r, 1).Formula = f"=ROW()-{HEADER_ROW}"
        ws.Range(f"F{DATA_START}:F{DATA_END}").NumberFormat = "0.0"
        ws.Range(f"G{DATA_START}:H{DATA_END}").NumberFormat = "yyyy/mm/dd"
        ws.Range(f"B{DATA_START}:B{DATA_END}").NumberFormat = "@"
        # 入力規則：状態（リスト直書き）、担当（別シートの範囲）
        v = ws.Range(f"I{DATA_START}:I{DATA_END}").Validation
        v.Delete()
        v.Add(Type=3, AlertStyle=1, Operator=1, Formula1="未着手,対応中,完了")
        v = ws.Range(f"E{DATA_START}:E{DATA_END}").Validation
        v.Delete()
        v.Add(Type=3, AlertStyle=1, Operator=1, Formula1="=リスト!$A$1:$A$4")
        # 合計行（明細の外）
        ws.Cells(DATA_END + 1, 4).Value = "合計"
        ws.Cells(DATA_END + 1, 6).Formula = f"=SUM(F{DATA_START}:F{DATA_END})"
        ws.Range(f"A{DATA_END + 1}:J{DATA_END + 1}").Borders.LineStyle = 1
        ws.Range(f"A{DATA_END + 1}:J{DATA_END + 1}").Font.Bold = True
        # 条件付き書式：状態が完了なら灰色
        fc = ws.Range(f"I{DATA_START}:I{DATA_END}").FormatConditions.Add(Type=1, Operator=3, Formula1='="完了"')
        fc.Interior.Color = 0xD9D9D9
        # 名前定義・列幅・印刷設定
        wb.Names.Add(Name="WBS範囲", RefersTo=f"=WBS!$A${HEADER_ROW}:$J${DATA_END}")
        for col, w in zip("ABCDEFGHIJ", (5, 9, 16, 30, 16, 9, 11, 11, 9, 20)):
            ws.Columns(col).ColumnWidth = w
        ws.PageSetup.Orientation = 2
        ws.PageSetup.PrintArea = f"$A$1:$J${DATA_END + 1}"
        ws.PageSetup.PrintTitleRows = f"${HEADER_ROW}:${HEADER_ROW}"
        # VBA：標準モジュールと Workbook_Open
        mod = wb.VBProject.VBComponents.Add(1)
        mod.Name = "Module1"
        mod.CodeModule.AddFromString(VBA_MODULE.split("\n", 1)[1])
        wb.VBProject.VBComponents("ThisWorkbook").CodeModule.AddFromString(VBA_WORKBOOK_OPEN)
        # ボタン（フォームコントロール）と図形
        left, top = ws.Range("L2").Left, ws.Range("L2").Top
        btn = ws.Buttons().Add(left, top, 110, 24)
        btn.OnAction = "Module1.RenumberRows"
        btn.Caption = "番号振り直し"
        shp = ws.Shapes.AddShape(1, left, top + 40, 110, 30)
        shp.TextFrame2.TextRange.Text = "記入例は削除"
        ws.Activate()
        ws.Range("A1").Select()

        wb.SaveAs(str(path), FileFormat=52)
        # 保護付きの版
        for col in ("B", "C", "D", "F", "G", "H", "I", "J"):
            ws.Range(f"{col}{DATA_START}:{col}{DATA_END}").Locked = False
        ws.Protect(DrawingObjects=False, Contents=True, Scenarios=False)
        wb.SaveAs(str(protected_path), FileFormat=52)
        wb.Close(SaveChanges=False)
    finally:
        quit_and_wait(xl)
        del xl
        pythoncom.CoUninitialize()


# ---------------------------------------------------------------------------
# 「大分類・中分類・作業項目」のひな形と、「管理単位（大）〜（小）」の実用WBS（Excel 不要）
# ---------------------------------------------------------------------------

UNIT_MASTER = [   # (大分類, 中分類, 作業項目, 担当)
    ("基本設計", "画面設計", "画面一覧の作成", "ベンダー"),
    ("基本設計", "画面設計", "画面レイアウトの作成", "ベンダー"),
    ("基本設計", "帳票設計", "帳票一覧の作成", "ベンダー"),
    ("基本設計", "帳票設計", "帳票レイアウトの作成", "医事課"),
    ("詳細設計", "連携設計", "電子カルテ連携仕様の作成", "ベンダー"),
    ("詳細設計", "連携設計", "PACS連携仕様の作成", "ベンダー"),
    ("詳細設計", "マスタ設計", "点数マスタ移行設計", "ベンダー"),
]
UNIT_HEADERS = ["管理単位（大）", "管理単位（中１）", "管理単位（中２）", "管理単位（中３）", "管理単位（小）", "担当", "備考"]
UNIT_CASES = {
    "units_A": {"skip": [], "rename": {}},
    "units_B": {"skip": ["PACS連携仕様の作成"], "rename": {"画面一覧の作成": "画面一覧作成"}},
}


def _grouped(rows):
    """分類はグループの先頭行にだけ書く（実際の WBS によくある書き方）。"""
    prev1 = prev2 = None
    for c1, c2, *rest in rows:
        yield (c1 if c1 != prev1 else None, c2 if (c1, c2) != (prev1, prev2) else None, *rest)
        prev1, prev2 = c1, c2


def build_unit_fixtures(here: Path = HERE) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WBSひな形"
    ws["A1"] = "設計WBS ひな形"
    for i, h in enumerate(["大分類", "中分類", "作業項目", "担当", "備考"], 1):
        c = ws.cell(2, i, h)
        c.font, c.fill, c.border = Font(bold=True), HEAD_FILL, BOX
    for r, row in enumerate(_grouped(UNIT_MASTER), 3):
        for i, v in enumerate(row, 1):
            ws.cell(r, i, v).border = BOX
    wb.save(here / "master_units.xlsx")

    def practical(path: Path, rows, n_blank: int = 0, title: str = ""):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "WBS"
        ws["A1"] = title or "医事会計システム更新 設計WBS"
        for i, h in enumerate(UNIT_HEADERS, 1):
            c = ws.cell(3, i, h)
            c.font, c.fill, c.border = Font(bold=True), HEAD_FILL, BOX
        r = 4
        for c1, c2, name, owner in _grouped(rows):
            for i, v in enumerate(["医事会計更新", "設計", c1, c2, name, owner, None], 1):
                ws.cell(r, i, v).border = BOX
            r += 1
        for _ in range(n_blank):
            for i in range(1, len(UNIT_HEADERS) + 1):
                ws.cell(r, i).border = BOX
            r += 1
        for col, w in zip("ABCDEFG", (14, 12, 14, 14, 26, 12, 16)):
            ws.column_dimensions[col].width = w
        wb.save(path)

    for name, spec in UNIT_CASES.items():
        rows = [(c1, c2, spec["rename"].get(n, n), o) for c1, c2, n, o in UNIT_MASTER if n not in spec["skip"]]
        practical(here / f"{name}.xlsx", rows)
    # 出力テンプレート：見出しと、罫線付きの空き行 15 行
    practical(here / "template_units.xlsx", [], n_blank=15, title="医事会計システム更新 WBS（管理単位）")


def scrub_metadata(path: Path) -> None:
    """ファイルのプロパティから作成者・最終更新者（Office のユーザー名）を消す。VBA など他のパーツは変えない。"""
    import re as _re
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        parts = {i.filename: z.read(i.filename) for i in infos}
    core = parts.get("docProps/core.xml")
    if core is None:
        return
    core = _re.sub(rb"<dc:creator>.*?</dc:creator>", b"<dc:creator></dc:creator>", core)
    core = _re.sub(rb"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", b"<cp:lastModifiedBy></cp:lastModifiedBy>", core)
    parts["docProps/core.xml"] = core
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w") as z:
        for i in infos:
            z.writestr(i, parts[i.filename], compress_type=i.compress_type)
    tmp.replace(path)


def main() -> int:
    build_master(HERE / "master.xlsx")
    for name, spec, title in (("case_A", CASE_A, "A病院 医事会計更新"), ("case_B", CASE_B, "B病院 医事会計更新"),
                              ("case_C", CASE_C, "C病院 医事会計更新")):
        build_case(HERE / f"{name}.xlsx", spec, title)
    build_unit_fixtures()
    if "--no-excel" in sys.argv:          # Excel を使わない素材だけ作る
        for p in sorted(HERE.glob("*units*.xlsx")):
            scrub_metadata(p)
        return 0
    version = excel_version()
    print("Excel", version)
    with TrustVBOM(version):
        build_template(HERE / "wbs_template.xlsm", HERE / "wbs_template_protected.xlsm")
    for p in sorted(HERE.glob("*.xls*")):
        scrub_metadata(p)          # 公開しても個人名が残らないように
    for p in sorted(HERE.glob("*.xls*")):
        print(p.name, p.stat().st_size)
    return 0


if __name__ == "__main__":
    sys.exit(main())
